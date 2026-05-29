from __future__ import annotations

import csv
import io
import uuid
from typing import Any

from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.phone_validation import phone_validation_to_storage, validate_phone_number
from models import CampaignExternalLead, User

NAME_HEADERS = frozenset({"nome", "name", "cliente", "nome_cliente"})
PHONE_HEADERS = frozenset({"telefone", "phone", "whatsapp", "celular", "fone", "numero", "número"})


def try_normalize_phone(raw: str) -> str | None:
    result = validate_phone_number(raw)
    if not result["valid"]:
        return None
    return phone_validation_to_storage(result["formatted"])


def _normalize_header(value: str) -> str:
    return (value or "").strip().lower().replace(" ", "_")


def _pick_column(fieldnames: list[str] | None, candidates: frozenset[str]) -> str | None:
    if not fieldnames:
        return None
    normalized = {_normalize_header(h): h for h in fieldnames if h}
    for key in candidates:
        if key in normalized:
            return normalized[key]
    return None


def _rows_from_csv(text: str) -> tuple[list[str], list[dict[str, str]]]:
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Arquivo sem cabeçalho.")
    rows: list[dict[str, str]] = []
    for row in reader:
        rows.append({k: (v or "") for k, v in row.items()})
    return list(reader.fieldnames), rows


def _rows_from_xlsx(file_bytes: bytes) -> tuple[list[str], list[dict[str, str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Importação Excel (.xlsx) indisponível no servidor. Use CSV ou contate o suporte.",
        ) from exc
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Planilha vazia.")
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Planilha vazia.")
    fieldnames = [str(c or "").strip() for c in header_row]
    rows: list[dict[str, str]] = []
    for row in rows_iter:
        if not any(row):
            continue
        item = {}
        for i, key in enumerate(fieldnames):
            if not key:
                continue
            val = row[i] if i < len(row) else ""
            item[key] = "" if val is None else str(val).strip()
        rows.append(item)
    return fieldnames, rows


def parse_leads_file(*, file_bytes: bytes, filename: str | None) -> tuple[list[str], list[dict[str, str]]]:
    lower = (filename or "").lower()
    if lower.endswith(".csv"):
        text = file_bytes.decode("utf-8-sig")
        return _rows_from_csv(text)
    if lower.endswith(".xlsx"):
        return _rows_from_xlsx(file_bytes)
    if lower.endswith(".xls"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Formato .xls não suportado. Salve como .xlsx ou .csv.",
        )
    try:
        return _rows_from_csv(file_bytes.decode("utf-8-sig"))
    except UnicodeDecodeError:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Formato não reconhecido. Use .csv ou .xlsx com colunas Nome e Telefone.",
        ) from None


def build_import_template_xlsx() -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Geração de modelo indisponível no servidor.",
        ) from exc

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    headers = ["Nome", "Telefone"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(["Maria Silva", "11999887766"])
    ws.append(["João Santos", "+55 21 98765-4321"])
    ws.append(["Empresa ABC Ltda", "5531988776655"])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    note = wb.create_sheet("Instruções")
    note.append(["Coluna", "Descrição"])
    note.append(["Nome", "Nome do contato (obrigatório)"])
    note.append(["Telefone", "DDD + número. Ex.: 11999887766 ou +55 11 99988-7766"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def validate_leads_from_file(
    *,
    file_bytes: bytes,
    filename: str | None,
) -> dict[str, Any]:
    """Processa o arquivo inteiro sem gravar no banco."""
    fieldnames, rows = parse_leads_file(file_bytes=file_bytes, filename=filename)
    name_col = _pick_column(fieldnames, NAME_HEADERS)
    phone_col = _pick_column(fieldnames, PHONE_HEADERS)
    if not name_col or not phone_col:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="O arquivo deve conter colunas 'Nome' e 'Telefone' (ou equivalentes).",
        )

    valid_rows: list[dict[str, Any]] = []
    invalid_rows: list[dict[str, Any]] = []
    skipped_empty = 0
    skipped_duplicate = 0
    seen_phones: set[str] = set()

    for line_no, row in enumerate(rows, start=2):
        name = (row.get(name_col) or "").strip()
        phone_raw = (row.get(phone_col) or "").strip()
        if not name and not phone_raw:
            skipped_empty += 1
            continue
        if not name:
            invalid_rows.append(
                {
                    "line_no": line_no,
                    "name": "",
                    "phone_raw": phone_raw,
                    "error": "Nome obrigatório.",
                }
            )
            continue

        validation = validate_phone_number(phone_raw)
        if not validation["valid"]:
            invalid_rows.append(
                {
                    "line_no": line_no,
                    "name": name,
                    "phone_raw": phone_raw,
                    "error": validation["error"] or "Telefone inválido.",
                }
            )
            continue

        phone_storage = phone_validation_to_storage(validation["formatted"])
        if phone_storage in seen_phones:
            skipped_duplicate += 1
            continue
        seen_phones.add(phone_storage)
        valid_rows.append(
            {
                "line_no": line_no,
                "name": name[:160],
                "phone": phone_storage,
                "formatted": validation["formatted"],
                "whatsapp_preview": (
                    f"{phone_storage[:4]}…{phone_storage[-4:]}" if len(phone_storage) > 8 else phone_storage
                ),
            }
        )

    total_parsed = len(valid_rows) + len(invalid_rows) + skipped_empty + skipped_duplicate
    summary = (
        f"Validação concluída: {len(valid_rows)} número(s) prontos"
        + (f", {len(invalid_rows)} com formato inválido" if invalid_rows else "")
        + (f", {skipped_duplicate} duplicado(s) ignorado(s)" if skipped_duplicate else "")
        + "."
    )
    return {
        "valid_rows": valid_rows,
        "invalid_rows": invalid_rows,
        "valid_count": len(valid_rows),
        "invalid_count": len(invalid_rows),
        "skipped_empty": skipped_empty,
        "skipped_duplicate": skipped_duplicate,
        "requires_review": len(invalid_rows) > 0,
        "source_filename": filename,
        "total_rows_parsed": total_parsed,
        "validation_message": summary,
    }


def confirm_external_leads(
    db: Session,
    *,
    tenant_id: int,
    user: User,
    leads: list[dict[str, str]],
    source_filename: str | None = None,
    discarded_invalid_count: int = 0,
) -> dict[str, Any]:
    """Persiste leads já validados/corrigidos pelo usuário."""
    if not leads:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Nenhum contato para importar.",
        )

    batch_id = str(uuid.uuid4())
    imported: list[dict[str, Any]] = []
    discarded = 0
    seen_phones: set[str] = set()

    for item in leads:
        name = (item.get("name") or "").strip()
        phone_input = (item.get("phone") or item.get("phone_raw") or "").strip()
        if not name or not phone_input:
            discarded += 1
            continue
        validation = validate_phone_number(phone_input)
        if not validation["valid"]:
            discarded += 1
            continue
        phone = phone_validation_to_storage(validation["formatted"])
        if phone in seen_phones:
            discarded += 1
            continue
        seen_phones.add(phone)
        lead = CampaignExternalLead(
            tenant_id=tenant_id,
            import_batch_id=batch_id,
            name=name[:160],
            phone=phone,
            source_filename=(source_filename or "import")[:180],
            created_by_user_id=user.id,
        )
        db.add(lead)
        db.flush()
        imported.append(
            {
                "id": lead.id,
                "name": lead.name,
                "phone": lead.phone,
                "formatted": validation["formatted"],
                "whatsapp_preview": (
                    f"{phone[:4]}…{phone[-4:]}" if len(phone) > 8 else phone
                ),
            }
        )

    if not imported:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Nenhum número válido para importar após a confirmação.",
        )

    db.commit()
    total_discarded = discarded + max(0, discarded_invalid_count)
    return {
        "import_batch_id": batch_id,
        "imported_count": len(imported),
        "discarded_count": total_discarded,
        "skipped_count": 0,
        "errors": [],
        "leads": imported,
        "summary_message": (
            f"Importação pronta! {len(imported)} número(s) validados e prontos para a campanha. "
            f"{total_discarded} número(s) foram descartados por formato inválido."
        ),
    }


def import_external_leads(
    db: Session,
    *,
    tenant_id: int,
    user: User,
    file_bytes: bytes,
    source_filename: str | None,
) -> dict[str, Any]:
    """Fluxo direto (legado): valida e grava em uma etapa."""
    preview = validate_leads_from_file(file_bytes=file_bytes, filename=source_filename)
    if preview["requires_review"]:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={
                "message": f"Identificamos {preview['invalid_count']} número(s) fora do formato. Revise a lista antes de importar.",
                "invalid_count": preview["invalid_count"],
                "valid_count": preview["valid_count"],
                "invalid_rows": preview["invalid_rows"][:100],
                "valid_rows": preview["valid_rows"],
            },
        )
    leads_payload = [{"name": r["name"], "phone": r["phone"]} for r in preview["valid_rows"]]
    result = confirm_external_leads(
        db,
        tenant_id=tenant_id,
        user=user,
        leads=leads_payload,
        source_filename=source_filename,
    )
    result["skipped_count"] = preview["skipped_duplicate"] + preview["skipped_empty"]
    result["errors"] = []
    return result
